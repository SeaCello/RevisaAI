# imports for create_flashcards
from django.shortcuts import render
from .models import UserFlashcard
from .forms import CreateCardForm
from gpt.api import generate_flashcards
from gpt.utils import extract_text_from_pdf
import chardet
import docx
from django.contrib.auth.decorators import login_required

# imports for dowload_pdf
from django.http import HttpResponse
from fpdf import FPDF

def create_flashcards(request):
    """
    Processa o arquivo enviado pelo usuário e gera flashcards a partir do seu conteúdo.
    Também permite ao usuário editar os flashcards gerados.
    """
    
    flashcards = []
    form = CreateCardForm()

    if request.method == 'POST':
        form = CreateCardForm(request.POST, request.FILES)
        
        if form.is_valid():
            try:
                file = request.FILES['file'] 
                #process PDF
                if file.name.lower().endswith('.pdf'):
                    text = extract_text_from_pdf(file)[:6000]  
                
                #process DOCX
                elif file.name.lower().endswith('.docx'):
                    doc = docx.Document(file)
                    text = '\n'.join([p.text for p in doc.paragraphs])[:6000]
                
                #process TXT
                else:
                    encoding = chardet.detect(file.read())['encoding']
                    file.seek(0)
                    text = file.read().decode(encoding)[:6000]

                
                flashcards = generate_flashcards(text)
                request.session['flashcards'] = flashcards 

                for flashcard in flashcards: # Save flashcards in database
                    UserFlashcard.objects.create(
                        user=request.user,
                        content=flashcard
                    )

            except Exception as e:
                form.add_error(None, f"Erro: {str(e)}")
    
        else:
            #Edit flashcards
            flashcards = request.POST.getlist('flashcards')
            request.session['flashcards'] = flashcards

    return render(request, 'create_flashcards.html', {
        'form': form,
        'flashcards': flashcards
    })
    


@login_required
def user_flashcards_home(request):
    """Página inicial do usuário para gestão de flashcards"""
    return render(
        request, 
        'home_user.html',
    )
    
    
def meus_flashcards(request):
    """Exibe os flashcards salvos pelo usuário"""
    user_flashcards = UserFlashcard.objects.filter(user=request.user)
    return render(request, 'meus_flashcards.html', {
        'flashcards': user_flashcards
    })

    
def download_pdf(request):
    """
    Recupera os flashcards armazenados na sessão do usuário e gera um documento PDF para download.
    
    O método verifica se existem flashcards na sessão. Se existirem, cria um PDF
    utilizando a biblioteca FPDF, adicionando um título e listando os flashcards.
    Em seguida, retorna o PDF como resposta HTTP para download.
    """
    
    flashcards = request.session.get('flashcards', [])
    
    if not flashcards:
        return HttpResponse("Nenhum flashcard encontrado.", content_type='text/plain')

    try:
        from io import BytesIO 
        # o pdf corrompe se não for feito isso
        buffer_pdf = BytesIO() 
        
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        
        pdf.set_font("Arial", 'B', 16)
        pdf.cell(0, 10, txt="Flashcards Gerados", ln=1, align='C')
        pdf.ln(10)

        pdf.set_font("Arial", size=12)
        for idx, card in enumerate(flashcards, 1):
            texto = f"{idx}. {card}"
            safe_text = texto.encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(0, 10, txt=safe_text)
            pdf.ln(8)

        pdf_bytes = pdf.output(dest='S').encode('latin-1')  # TRANSFORMA O PDF EM BYTES NO FORMATO STRING
        buffer_pdf.write(pdf_bytes)
        buffer_pdf.seek(0)
        
        response = HttpResponse(buffer_pdf.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="flashcards.pdf"'
        return response
    
    except Exception as e:
        return HttpResponse(f"Erro na geração do PDF: {str(e)}", status=500)



